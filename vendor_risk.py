"""CLI for the third-party vendor risk assessment framework."""
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from src.risk_engine import portfolio_metrics, score_assessment


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Automate a third-party vendor risk assessment.")
    parser.add_argument("--vendors", default="data/vendor_profiles.csv")
    parser.add_argument("--questions", default="data/question_bank.csv")
    parser.add_argument("--answers", default="data/vendor_answers.csv")
    parser.add_argument("--out-dir", default="reports")
    return parser


def write_excel(register: pd.DataFrame, findings: pd.DataFrame, output: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Risk Register"
    header_fill = PatternFill("solid", fgColor="173B3F")
    header_font = Font(color="FFFFFF", bold=True)
    for row in [register.columns.tolist(), *register.astype(object).where(pd.notna(register), "").values.tolist()]:
        summary.append(row)
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    summary.column_dimensions["B"].width = 24
    summary.column_dimensions["C"].width = 24
    summary.column_dimensions["I"].width = 18
    summary.column_dimensions["J"].width = 18

    detail = workbook.create_sheet("Assessment Findings")
    for row in [findings.columns.tolist(), *findings.astype(object).where(pd.notna(findings), "").values.tolist()]:
        detail.append(row)
    for cell in detail[1]:
        cell.fill = header_fill
        cell.font = header_font
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    for column in "ABCDEFGHIJKLM":
        detail.column_dimensions[column].width = 20

    workbook.save(output)


def write_pdf(register: pd.DataFrame, findings: pd.DataFrame, output: Path) -> None:
    metrics = portfolio_metrics(register)
    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(str(output), pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    story = [Paragraph("Third-Party Vendor Risk Assessment Report", styles["Title"]), Paragraph("Fictional portfolio demonstration — NIST SP 800-161 Rev. 1 methodology lens", styles["Normal"]), Spacer(1, 6 * mm)]
    metric_data = [["Vendors", "High/Critical", "Average residual", "Evidence coverage"], [metrics["vendors"], metrics["critical_high"], f"{metrics['average_residual_score']:.1f}/100", f"{metrics['evidence_coverage']:.1f}%"]]
    metric_table = Table(metric_data, colWidths=[35 * mm] * 4)
    metric_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173B3F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#B8C9C7")), ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#EDF3F1"))]))
    story.append(metric_table)
    story.append(Spacer(1, 7 * mm))
    story.append(Paragraph("Risk register", styles["Heading2"]))
    rows = [["Vendor", "Service", "Inherent", "Effectiveness", "Residual", "Tier", "Gaps", "Review"]]
    for _, row in register.iterrows():
        rows.append([row["vendor_name"], row["service_type"], f"{row['inherent_score']:.1f}", f"{row['control_effectiveness_percent']:.1f}%", f"{row['residual_score']:.1f}", row["risk_tier"], row["open_gaps"], row["recommended_review"]])
    table = Table(rows, repeatRows=1, colWidths=[35 * mm, 43 * mm, 20 * mm, 25 * mm, 20 * mm, 22 * mm, 15 * mm, 25 * mm])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173B3F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#B8C9C7")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7F6")])]))
    story.append(table)
    story.append(Spacer(1, 7 * mm))
    story.append(Paragraph("Method note: scores are explainable prioritization heuristics for fictional data. They are not an audit opinion, certification, or legal risk determination. Review the CSV/Excel detail for question-level remediation ownership.", styles["Normal"]))
    document.build(story)


def main() -> None:
    args = build_parser().parse_args()
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    vendors = pd.read_csv(args.vendors, keep_default_na=False)
    questions = pd.read_csv(args.questions, keep_default_na=False)
    answers = pd.read_csv(args.answers, keep_default_na=False)
    register, findings = score_assessment(vendors, questions, answers)
    register.to_csv(out_dir / "vendor_risk_register.csv", index=False)
    findings.to_csv(out_dir / "vendor_assessment_findings.csv", index=False)
    write_excel(register, findings, out_dir / "vendor_risk_assessment.xlsx")
    write_pdf(register, findings, out_dir / "vendor_risk_assessment.pdf")
    print(f"Generated reports in {out_dir.resolve()}")
    print(portfolio_metrics(register))


if __name__ == "__main__":
    main()
